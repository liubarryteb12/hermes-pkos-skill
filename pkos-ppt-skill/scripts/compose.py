#!/usr/bin/env python3
"""
compose.py — pkos.exit.ppt.compose v2.0 编排层

2026-08-31 用户裁定：PPT 出口转向**原生可编辑 PPTX**（吸收 ppt-master 设计），
v0.2 出图制裁定作废。gptimage2 降级为可选插图通道（--images）。

流水线（Plan → Do → Check）：
  RT-*(exit=ppt) + POL-* → deck_spec.build_deck_spec → design_spec.json
  → build_pptx.build → <route-id>.pptx（原生文本框/形状，PowerPoint 可编辑）
  → 校验门（重开 pptx 数页 + hash）→ manifest.json（可复现）

契约保留：比例三选一必问 / 三态失败 / manifest 完整 / provider 铁律 /
degraded_success（插图失败但 deck 与 prompt 清单完整交付）。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import time
from pathlib import Path
from typing import Any

# ── 路径解析 ────────────────────────────────────────────────────────────────

def _find_workspace_root() -> Path:
    """平铺布局：pkos-ppt-skill/scripts 的上两级即套件根。"""
    return Path(__file__).resolve().parents[2]

WORKSPACE_ROOT = _find_workspace_root()
SKILL_DIR = WORKSPACE_ROOT / "pkos-ppt-skill"

sys.path.insert(0, str(SKILL_DIR / "scripts"))
_img_api_path = SKILL_DIR / "shared" / "image-api"
if str(_img_api_path) not in sys.path:
    sys.path.insert(0, str(_img_api_path))

import deck_spec  # noqa: E402
import build_pptx  # noqa: E402

try:
    import client as img_client  # type: ignore  # noqa: E402
except ImportError:
    img_client = None  # 插图通道缺失不影响原生 deck

# ── 比例映射（v0 锁死继承：三选一必问）─────────────────────────────────────

RATIO_SIZES: dict[str, str] = {
    "16:9": "1536x1024",   # 插图槽位用（image-size-spec 白名单）
    "4:3": "1536x1024",
    "3:4": "1024x1536",
}
DEFAULT_RATIO = "16:9"

# conversion_type 承接子集：对齐 router v3.3 兼容性矩阵（ppt 行 4 值全合法）。
# 旧实现只放 2 值且与 router 漂移，2026-08-31 修正。
PPT_ACCEPTS = {"实战操作指南", "wiki百科条目", "避坑风险清单", "学习路径"}


# ── YAML 解析（最小实现，避免 pyyaml 依赖）──────────────────────────────────

def yaml_safe_load(text: str) -> dict:
    result: dict[str, Any] = {}
    current_list_key: str | None = None
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        if stripped.startswith("- "):
            if current_list_key:
                val = stripped[2:].strip().strip('"').strip("'")
                if isinstance(result.get(current_list_key), list):
                    result[current_list_key].append(val)
                else:
                    result[current_list_key] = [val]
            continue
        m = re.match(r"^(\S+):\s*(.*)$", stripped)
        if m:
            key, val = m.group(1), m.group(2).strip()
            current_list_key = key
            if val == "":
                result[key] = []
            elif val.startswith("[") and val.endswith("]"):
                result[key] = [v.strip().strip('"').strip("'")
                               for v in val[1:-1].split(",") if v.strip()]
            else:
                val = val.strip('"').strip("'")
                if val.lower() == "true":
                    result[key] = True
                elif val.lower() == "false":
                    result[key] = False
                elif val.isdigit():
                    result[key] = int(val)
                else:
                    result[key] = val
    return result


def load_route(route_path: Path) -> dict:
    if not route_path.exists():
        raise FileNotFoundError(f"路由单不存在: {route_path}")
    with open(route_path, "r", encoding="utf-8-sig") as f:
        return yaml_safe_load(f.read())


def validate_route(route: dict) -> list[str]:
    rejections: list[str] = []
    if route.get("exit") != "ppt":
        rejections.append(f"本单元只消费 exit=ppt，路由单 exit={route.get('exit')}")
    conv = route.get("conversion_type", "")
    if conv and conv not in PPT_ACCEPTS:
        rejections.append(f"conversion_type '{conv}' 不在 ppt 承接子集 {sorted(PPT_ACCEPTS)}")
    if not route.get("source_entry"):
        rejections.append("路由单缺少 source_entry")
    return rejections


# ── POL 源解析 ──────────────────────────────────────────────────────────────

def resolve_source(route: dict) -> Path | None:
    source = route.get("source_entry", "")
    m = re.search(r"\[\[(.+?)\]\]", source)
    if not m:
        return None
    name = m.group(1)
    analysis_dir = WORKSPACE_ROOT / "_PKOS" / "analysis"
    if not analysis_dir.exists():
        return None
    candidates = list(analysis_dir.glob(f"*{name}*"))
    if candidates:
        return max(candidates, key=lambda p: p.stat().st_mtime)
    exact = analysis_dir / f"{name}.md"
    return exact if exact.exists() else None


def parse_pol_content(path: Path) -> dict:
    text = path.read_text(encoding="utf-8-sig")
    if text.startswith("---"):
        parts = text.split("---", 2)
        fm_lines = parts[1].splitlines() if len(parts) > 1 else []
        body = parts[2] if len(parts) > 2 else ""
    else:
        fm_lines = []
        body = text
    fm: dict[str, str] = {}
    for line in fm_lines:
        m = re.match(r"^(\w[\w-]*):\s*(.*)$", line.strip())
        if m:
            fm[m.group(1)] = m.group(2).strip().strip('"')
    # 正文首个 H1 = deck 标题（frontmatter title 常是 POL 机器 ID，不可用作标题）
    h1 = re.search(r"^# (.+)$", body, re.M)
    deck_title = h1.group(1).strip() if h1 else ""
    # 元信息编号不入标题（用户规则：EP/章节号前缀剥离）
    deck_title = re.sub(r"^\d+(?:[-.]\d+)*\s*", "", deck_title)
    # "POL-xxx —— 真实标题" 形态：取破折号后段
    if "——" in deck_title:
        deck_title = deck_title.split("——")[-1].strip()
    deck_title = re.sub(r"\s*·\s*润色净本\s*$", "", deck_title)
    sections = []
    for sec in re.split(r"^## ", body, flags=re.M):
        sec = sec.strip()
        if not sec:
            continue
        lines = sec.splitlines()
        title = lines[0].strip() if lines else ""
        sec_body = "\n".join(lines[1:]).strip() if len(lines) > 1 else ""
        # H3 子节展开为独立页（吸收 ppt-master：一页一个主张，不做大杂烩页）
        subs = list(re.finditer(r"^### (.+)$", sec_body, re.M))
        if len(subs) >= 2:
            for j, sm in enumerate(subs):
                end = subs[j + 1].start() if j + 1 < len(subs) else len(sec_body)
                sub_body = sec_body[sm.end():end].strip()
                st = sm.group(1).strip()
                st = re.sub(r"^[一二三四五六七八九十]+、", "", st)  # 去中文序号
                sections.append({"title": st, "body": sub_body})
            pre = sec_body[: subs[0].start()].strip()
            if pre:
                sections.insert(len(sections) - len(subs), {"title": title, "body": pre})
        else:
            sections.append({"title": title, "body": sec_body})
    # 段落式净本兜底：无 ## 分节的 POL（digest 型），按段落分组分页
    if not any(not s["title"].startswith("#") for s in sections):
        main = re.split(r"\n-{3,}\n", body)[0]  # 去掉尾部评分卡/改动追溯元信息段
        paras = [p.strip() for p in re.split(r"\n\s*\n", main)
                 if p.strip() and not p.strip().startswith("#")]
        paras = [p for p in paras if "五维评分卡" not in p and "改动追溯" not in p]
        if paras:
            cn = "一二三四五六七八九十"
            n_groups = min(max(len(paras) // 2, 2), 5)
            chunk = -(-len(paras) // n_groups)
            for gi in range(0, len(paras), chunk):
                g = gi // chunk
                sections.append({"title": f"要点{cn[g] if g < 10 else g + 1}",
                                 "body": "\n\n".join(paras[gi: gi + chunk])})
    return {"frontmatter": fm, "deck_title": deck_title,
            "sections": sections, "raw": body}


# ── 美学主题 ────────────────────────────────────────────────────────────────

def load_aesthetics() -> dict:
    p = SKILL_DIR / "themes" / "aesthetics.json"
    if p.exists():
        with open(p, "r", encoding="utf-8-sig") as f:
            return json.load(f)
    return {"themes": {}, "fallback_theme": {}, "typography_scale": {}}


def resolve_theme(route: dict, aesthetics: dict) -> str:
    style = route.get("style_adapter") or route.get("style_theme")
    if style and str(style).strip() and str(style) != "null" and str(style) in aesthetics.get("themes", {}):
        return str(style)
    return "paper-ink"


# ── 比例确认（v0 锁死：必问，不得默认）──────────────────────────────────────

class RatioRequiredError(Exception):
    pass


def ask_ratio(ratio_input: str | None, auto_mode: bool) -> tuple[str, str | None]:
    if ratio_input and ratio_input in deck_spec.CANVAS:
        return ratio_input, None
    if auto_mode:
        return DEFAULT_RATIO, "auto_mode: 默认 16:9"
    raise RatioRequiredError(
        "生图比例未确认（v0 锁死铁律），三选一：\n"
        "  1. 16:9 横版宽屏（推荐，投屏最通用）\n"
        "  2. 4:3  横版标准\n"
        "  3. 3:4  竖版\n"
        "回复数字或比例（如 '16:9' / '竖版'）。"
    )


# ── 插图槽位（可选，gptimage2）──────────────────────────────────────────────

def fill_image_slots(spec: dict, aesthetics: dict, out_dir: Path,
                     max_images: int = 3) -> tuple[list[dict], list[dict]]:
    """把最多 max_images 个内容页升级为 image-right 并生成插图。
    返回 (done[{slide,file,prompt}], failed[{slide,prompt,error}])。"""
    theme = aesthetics.get("themes", {}).get(spec["theme"], aesthetics.get("fallback_theme", {}))
    tokens = theme.get("prompt_tokens", "")
    size = RATIO_SIZES.get(spec["canvas"]["ratio"], "1536x1024")
    candidates = [s for s in spec["slides"]
                  if s["role"] in ("bullets", "two-column") and not s.get("emphasis")][::2]
    candidates = candidates[:max_images]
    done, failed = [], []
    if candidates and img_client is None:
        for s in candidates:
            failed.append({"slide": s["index"], "error": "image-api client 不可导入",
                           "prompt": ""})
        return done, failed
    cfg = None
    if candidates:
        try:
            cfg = img_client.load_config()
        except Exception as e:
            for s in candidates:
                failed.append({"slide": s["index"], "error": f"config 不可读: {e}", "prompt": ""})
            return done, failed
    for s in candidates:
        prompt = f"{tokens}, editorial illustration for a slide titled '{s['title']}', no text, no letters"
        try:
            r = img_client.call_generate(cfg=cfg, prompt=prompt, size=size, output_dir=out_dir)
            dest = out_dir / f"illustration-{s['index']:02d}.png"
            src = Path(r["file"])
            if src.resolve() != dest.resolve():
                if dest.exists():
                    dest.unlink()
                src.rename(dest)
            s["role"] = "image-right"
            s["image_slot"] = {"file": str(dest), "prompt": prompt,
                               "provider": r["provider"], "size": size}
            done.append({"slide": s["index"], "file": str(dest), "prompt": prompt,
                         "provider": r["provider"], "size": size})
        except Exception as e:
            failed.append({"slide": s["index"], "prompt": prompt, "error": str(e)})
    return done, failed


# ── 校验门（Check）──────────────────────────────────────────────────────────

def verify_pptx(pptx_path: Path, expected_slides: int) -> list[str]:
    """gate_2_integrity 的原生 deck 版：重开文件数页、非空、可解析。"""
    errs: list[str] = []
    if not pptx_path.exists() or pptx_path.stat().st_size < 10_000:
        errs.append(f"pptx 缺失或异常小: {pptx_path}")
        return errs
    try:
        from pptx import Presentation
        prs = Presentation(str(pptx_path))
        if len(prs.slides) != expected_slides:
            errs.append(f"页数不符: 期望 {expected_slides}，实际 {len(prs.slides)}")
    except Exception as e:
        errs.append(f"pptx 无法重新打开: {e}")
    return errs


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


# ── 主流程 ──────────────────────────────────────────────────────────────────

def compose(
    route_path: str,
    ratio: str | None = None,
    slides_count: int | None = None,
    auto_mode: bool = False,
    theme: str | None = None,
    with_images: bool = False,
    out_dir: str | None = None,
) -> dict:
    route_file = Path(route_path)
    if not route_file.is_absolute():
        route_file = WORKSPACE_ROOT / route_path

    route = load_route(route_file)
    route_id = route.get("route_id", "unknown")

    rejections = validate_route(route)
    if rejections:
        return {"failure_mode": "not_found", "route_id": route_id,
                "rejections": rejections, "status": "rejected"}

    pol_path = resolve_source(route)
    if not pol_path:
        return {"failure_mode": "not_found", "route_id": route_id,
                "rejections": [f"源条目不可达: {route.get('source_entry')}"],
                "status": "rejected"}

    pol_text = pol_path.read_text(encoding="utf-8-sig")
    fm_match = re.search(r"^status:\s*(\S+)", pol_text, re.M)
    if fm_match and fm_match.group(1) not in ("polished", "routed", "exported"):
        return {"failure_mode": "not_found", "route_id": route_id,
                "rejections": [f"源条目 status={fm_match.group(1)}，早于 polished"],
                "status": "rejected"}

    try:
        final_ratio, decision_note = ask_ratio(ratio, auto_mode)
    except RatioRequiredError as e:
        return {"failure_mode": "ambiguous", "route_id": route_id,
                "decision_card": str(e), "status": "blocked"}

    pol_data = parse_pol_content(pol_path)
    aesthetics = load_aesthetics()
    theme_id = theme if theme in aesthetics.get("themes", {}) else resolve_theme(route, aesthetics)

    try:
        spec = deck_spec.build_deck_spec(route, pol_data, final_ratio, slides_count,
                                         theme_id, aesthetics)
    except Exception as e:
        return {"failure_mode": "unavailable", "route_id": route_id,
                "errors": [f"design_spec 生成失败: {e}"], "status": "failed"}
    if len(spec["slides"]) < 2:
        return {"failure_mode": "unavailable", "route_id": route_id,
                "errors": ["POL 素材不足，无法成册（至少封面+1 内容页）"], "status": "failed"}
    for i, s in enumerate(spec["slides"], 1):
        s["index"] = i

    out = Path(out_dir) if out_dir else WORKSPACE_ROOT / "_PKOS" / "outputs" / f"{route_id}-deck"
    out.mkdir(parents=True, exist_ok=True)

    # 插图槽位（可选通道；失败只降级，不阻塞原生 deck）
    img_done, img_failed = ([], [])
    if with_images:
        img_done, img_failed = fill_image_slots(spec, aesthetics, out)

    # design_spec 落盘（可审计、可手改后重渲染）
    spec_path = out / "design_spec.json"
    spec_path.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")

    # 渲染
    pptx_path = out / f"{route_id}.pptx"
    theme_native = aesthetics.get("themes", {}).get(theme_id, aesthetics.get("fallback_theme", {}))["native"]

    def resolver(sl: dict):
        slot = sl.get("image_slot")
        return slot["file"] if slot and slot.get("file") else None

    try:
        build_info = build_pptx.build(spec, theme_native, pptx_path, image_resolver=resolver)
    except Exception as e:
        return {"failure_mode": "unavailable", "route_id": route_id,
                "errors": [f"pptx 渲染失败: {e}"], "status": "failed"}

    # 校验门
    verify_errs = verify_pptx(pptx_path, len(spec["slides"]))
    if verify_errs:
        return {"failure_mode": "unavailable", "route_id": route_id,
                "errors": verify_errs, "status": "failed"}

    # manifest（可复现：spec 输入 + 每页角色/要点/讲稿 + 插图 provider）
    manifest = {
        "route_id": route_id,
        "schema": "pkos-ppt-deck:2",
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "ratio": final_ratio,
        "theme": theme_id,
        "degraded": bool(img_failed),
        "decision_note": decision_note,
        "pptx": str(pptx_path),
        "pptx_sha256": sha256(pptx_path),
        "design_spec": str(spec_path),
        "source": {"route": str(route_file), "pol": str(pol_path),
                   "pol_sha256": sha256(pol_path)},
        "provider": {"deck": "python-pptx:native",
                     "illustrations": [d["provider"] for d in img_done] or "none"},
        "slides": [{
            "slide": s["index"], "role": s["role"], "title": s["title"],
            "bullets": s.get("bullets", []), "notes": s.get("notes"),
            "image": s.get("image_slot"),
        } for s in spec["slides"]],
        "illustrations_failed": img_failed,
    }
    manifest_path = out / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    status = "degraded_success" if img_failed else "success"
    result = {
        "route_id": route_id,
        "status": status,
        "total_slides": len(spec["slides"]),
        "pptx": str(pptx_path),
        "sha256": manifest["pptx_sha256"],
        "design_spec": str(spec_path),
        "manifest": str(manifest_path),
        "output_dir": str(out),
        "ratio": final_ratio,
        "theme": theme_id,
        "images_placed": len(img_done),
        "images_failed": len(img_failed),
        "degraded": bool(img_failed),
    }
    if img_failed:
        result["failure_mode"] = "degraded_success"
        result["errors"] = img_failed
    return result


# ── CLI ─────────────────────────────────────────────────────────────────────

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description="pkos.exit.ppt.compose v2.0 — 路由单+POL → 原生可编辑 PPTX")
    ap.add_argument("--route", required=True, help="路由单路径（相对或绝对）")
    ap.add_argument("--ratio", default=None, help="画布比例: 16:9 / 4:3 / 3:4（必问）")
    ap.add_argument("--slides", type=int, default=None, help="页数（null=按内容自动）")
    ap.add_argument("--auto", action="store_true", help="自主轮次：默认 16:9 留痕")
    ap.add_argument("--theme", default=None, help="强制美学主题 ID")
    ap.add_argument("--images", action="store_true",
                    help="启用 gptimage2 插图槽位（默认纯原生不出图）")
    ap.add_argument("--out", default=None, help="输出目录覆盖")
    ap.add_argument("--dry-run", action="store_true", help="只出 design_spec，不渲染")
    args = ap.parse_args(argv)

    try:
        if args.dry_run:
            route_file = Path(args.route)
            if not route_file.is_absolute():
                route_file = WORKSPACE_ROOT / route_file
            route = load_route(route_file)
            rej = validate_route(route)
            if rej:
                print(json.dumps({"status": "rejected", "rejections": rej},
                                 ensure_ascii=False, indent=2))
                return 1
            pol = resolve_source(route)
            if not pol:
                print(json.dumps({"status": "rejected",
                                  "error": f"源不可达: {route.get('source_entry')}"},
                                 ensure_ascii=False))
                return 1
            pol_data = parse_pol_content(pol)
            aesthetics = load_aesthetics()
            theme_id = args.theme or resolve_theme(route, aesthetics)
            ratio = args.ratio or DEFAULT_RATIO
            spec = deck_spec.build_deck_spec(route, pol_data, ratio, args.slides,
                                             theme_id, aesthetics)
            print(json.dumps(spec, ensure_ascii=False, indent=2))
            return 0

        result = compose(args.route, args.ratio, args.slides, args.auto,
                         theme=args.theme, with_images=args.images, out_dir=args.out)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result.get("status") in ("success", "degraded_success") else 1
    except FileNotFoundError as e:
        print(json.dumps({"error": str(e)}, ensure_ascii=False), file=sys.stderr)
        return 1
    except Exception as e:
        print(json.dumps({"error": str(e), "traceback": repr(e)}, ensure_ascii=False),
              file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
