#!/usr/bin/env python3
"""单元测试：pkos.exit.ppt.compose v2.0（原生 PPTX 线）

覆盖：
- YAML 最小解析器
- 路由单校验（not_found；承接子集对齐 router v3.3 矩阵 4 值）
- POL 源解析（H1 标题提取 + 编号剥离 + H3 子节展开 + 元信息节过滤）
- 比例确认（ambiguous vs auto）
- deck_spec（页型角色、要点拆页、两级层次素材、封底 takeaway、节奏）
- build_pptx（中西文空格、标点归一、主点识别）
- 端到端 smoke（spec → pptx → 重开验页数）
"""
import json
import sys
import tempfile
from pathlib import Path

_TMP_ROOT = Path(__file__).resolve().parent.parent.parent / ".tmp-test"
_TMP_ROOT.mkdir(exist_ok=True)

sys.stdout.reconfigure(encoding='utf-8')
_SCRIPTS_DIR = Path(__file__).resolve().parent.parent.parent / "scripts"
sys.path.insert(0, str(_SCRIPTS_DIR))

import importlib.util as _ilu
_spec = _ilu.spec_from_file_location("compose_sut", _SCRIPTS_DIR / "compose.py")
sut = _ilu.module_from_spec(_spec)
_spec.loader.exec_module(sut)
import deck_spec
import build_pptx


# ── YAML / 路由单 ───────────────────────────────────────────────────────────

def test_yaml_parser_basic():
    text = """route_id: RT-20260827-001
exit: ppt
conversion_type: 实战操作指南
rationale:
  - "第一条理由"
  - "第二条理由"
"""
    result = sut.yaml_safe_load(text)
    assert result["route_id"] == "RT-20260827-001"
    assert result["exit"] == "ppt"
    assert result["rationale"] == ["第一条理由", "第二条理由"]
    print("PASS: test_yaml_parser_basic")


def test_yaml_parser_brackets():
    text = 'source_entry: "[[POL-2026-08-26-05-01-37-G.md]]"\nexit: ppt\n'
    result = sut.yaml_safe_load(text)
    assert result["source_entry"] == "[[POL-2026-08-26-05-01-37-G.md]]"
    print("PASS: test_yaml_parser_brackets")


def test_validate_route_exit_not_ppt():
    rejections = sut.validate_route({"exit": "html", "source_entry": "[[POL-xxx.md]]"})
    assert any("exit=ppt" in r for r in rejections)
    print("PASS: test_validate_route_exit_not_ppt")


def test_validate_route_conversion_subset():
    """承接子集 = router v3.3 矩阵 ppt 行 4 值；漫画/小说必须拒。"""
    for ct in ("实战操作指南", "wiki百科条目", "避坑风险清单", "学习路径"):
        assert sut.validate_route({"exit": "ppt", "source_entry": "[[x]]",
                                   "conversion_type": ct}) == [], ct
    for ct in ("公众号漫画", "小说"):
        assert sut.validate_route({"exit": "ppt", "source_entry": "[[x]]",
                                   "conversion_type": ct}), ct
    print("PASS: test_validate_route_conversion_subset")


def test_validate_route_missing_source():
    rejections = sut.validate_route({"exit": "ppt"})
    assert any("source_entry" in r for r in rejections)
    print("PASS: test_validate_route_missing_source")


# ── 比例（v0 锁死：必问）────────────────────────────────────────────────────

def test_ratio_auto_mode():
    ratio, note = sut.ask_ratio(None, auto_mode=True)
    assert ratio == "16:9"
    assert note and "auto" in note
    print("PASS: test_ratio_auto_mode")


def test_ratio_explicit():
    for r in ("16:9", "4:3", "3:4"):
        ratio, note = sut.ask_ratio(r, auto_mode=False)
        assert ratio == r and note is None
    print("PASS: test_ratio_explicit")


def test_ratio_ambiguous():
    try:
        sut.ask_ratio(None, auto_mode=False)
        assert False, "应抛出 RatioRequiredError"
    except sut.RatioRequiredError:
        pass
    print("PASS: test_ratio_ambiguous")


# ── POL 解析 ────────────────────────────────────────────────────────────────

def test_parse_pol_content():
    pol = _TMP_ROOT / "parse-pol-test.md"
    pol.write_text("""---
title: POL-machine-id
status: polished
---

# 05-01-37 GraphQL实战：从理论到生产环境

## 钩子
开场段落。

## 正文
### 一、场景切入
内容A。

### 二、底层复盘
内容B。

## 五维评分卡
总分 41。
""", encoding="utf-8")
    data = sut.parse_pol_content(pol)
    assert data["frontmatter"]["status"] == "polished"
    # H1 提取 + EP 编号剥离
    assert data["deck_title"] == "GraphQL实战：从理论到生产环境"
    titles = [s["title"] for s in data["sections"]]
    # H3 子节展开为独立页
    assert "场景切入" in titles and "底层复盘" in titles
    print("PASS: test_parse_pol_content")


# ── deck_spec ───────────────────────────────────────────────────────────────

def _mk_pol():
    return {
        "frontmatter": {"title": "POL-id"},
        "deck_title": "测试主题：知识体系",
        "sections": [
            {"title": "钩子", "body": "第一段引入。\n\n第二段补充。"},
            {"title": "正文", "body": "\n".join(
                f"要点{i}：这是第{i}条说明文字，用于测试拆页逻辑。" for i in range(1, 9))},
            {"title": "CTA", "body": "> 金句引用测试，长度在八到一百二十之间。"},
        ],
    }


def test_deck_spec_basic():
    route = {"route_id": "RT-t", "topic_suggestion": "T", "audience": "开发者"}
    aesthetics = {"themes": {"paper-ink": {"native": {}}},
                  "typography_scale": {"cover_title": 2.6, "page_title": 1.6,
                                        "body": 1.0, "annotation": 0.78,
                                        "footnote": 0.62, "section_title": 2.0,
                                        "subtitle": 1.2}}
    spec = deck_spec.build_deck_spec(route, _mk_pol(), "16:9", None, "paper-ink", aesthetics)
    assert spec["schema"] == "pkos-ppt-deck-spec:2"
    assert spec["slides"][0]["role"] == "cover"
    assert spec["slides"][0]["title"] == "测试主题：知识体系"
    assert spec["slides"][-1]["role"] == "closing"
    # 8 条要点必须拆页（MAX_BULLETS=5）
    content_pages = [s for s in spec["slides"] if s["role"] in ("bullets", "two-column")]
    assert all(len(s["bullets"]) <= deck_spec.MAX_BULLETS * 2 for s in content_pages)
    # 画布
    assert spec["canvas"]["width_in"] == 13.333
    print("PASS: test_deck_spec_basic")


def test_deck_spec_quote_and_ratio():
    route = {"route_id": "RT-q"}
    aesthetics = {"themes": {}, "typography_scale": {}}
    spec = deck_spec.build_deck_spec(route, _mk_pol(), "3:4", None, "paper-ink", aesthetics)
    assert spec["canvas"]["height_in"] > spec["canvas"]["width_in"]
    assert spec["typography"]["body_pt"] == 15.0
    roles = [s["role"] for s in spec["slides"]]
    assert "quote" in roles  # CTA 节的 > 引用被识别为 quote 页
    print("PASS: test_deck_spec_quote_and_ratio")


def test_deck_spec_rejects_bad_ratio():
    try:
        deck_spec.build_deck_spec({"route_id": "x"}, _mk_pol(), "21:9", None,
                                  "paper-ink", {"themes": {}, "typography_scale": {}})
        assert False
    except ValueError:
        pass
    print("PASS: test_deck_spec_rejects_bad_ratio")


# ── build_pptx 排版归一 ─────────────────────────────────────────────────────

def test_cjk_spacing():
    assert build_pptx._cjk_punct("GraphQL实战") == "GraphQL 实战"
    assert build_pptx._cjk_punct("用BFF层") == "用 BFF 层"
    assert build_pptx._cjk_punct("保持不动.") == "保持不动。"
    assert build_pptx._cjk_punct("English only.") == "English only."
    print("PASS: test_cjk_spacing")


def test_main_bullet_detection():
    assert build_pptx._is_main_bullet("坑 1：N+1 查询")
    assert build_pptx._is_main_bullet("1. 本质是声明式查询")
    assert build_pptx._is_main_bullet("重来一次：用 BFF 渐进式落地")
    assert not build_pptx._is_main_bullet("列表页查一批作者，每个作者再查他的文章。")
    print("PASS: test_main_bullet_detection")


# ── 端到端 smoke：spec → pptx → 重开 ───────────────────────────────────────

def test_build_pptx_smoke():
    from pptx import Presentation
    aesthetics = {"themes": {}, "typography_scale": {"cover_title": 2.6, "page_title": 1.6,
                                                     "body": 1.0, "annotation": 0.78,
                                                     "footnote": 0.62, "section_title": 2.0,
                                                     "subtitle": 1.2}}
    spec = deck_spec.build_deck_spec({"route_id": "RT-smoke", "audience": "A"},
                                     _mk_pol(), "16:9", None, "paper-ink", aesthetics)
    theme = json.loads((sut.SKILL_DIR / "themes" / "aesthetics.json")
                       .read_text(encoding="utf-8-sig"))["themes"]["paper-ink"]["native"]
    out = _TMP_ROOT / "smoke.pptx"
    info = build_pptx.build(spec, theme, out)
    assert out.exists() and out.stat().st_size > 10_000
    prs = Presentation(str(out))
    assert len(prs.slides) == info["slides"] == len(spec["slides"])
    # 文本可编辑性：至少一页含原生文本框
    assert any(sh.has_text_frame and sh.text_frame.text.strip()
               for sh in prs.slides[0].shapes)
    print("PASS: test_build_pptx_smoke")


def test_verify_gate_catches_corrupt():
    bad = _TMP_ROOT / "bad.pptx"
    bad.write_bytes(b"not a pptx")
    errs = sut.verify_pptx(bad, 1)
    assert errs, "校验门必须拒绝损坏文件"
    print("PASS: test_verify_gate_catches_corrupt")




def test_theme_invalid_route_fail_loud():
    """v2.0.1: 路由单声明白名单外主题 → ThemeInvalidError，不静默回退"""
    route = {"style_theme": "not-a-theme"}
    try:
        sut.resolve_theme(route, {"themes": {"paper-ink": {}, "kan-shi": {}}})
        raise SystemExit("FAIL: 非法主题应 raise 而非静默回退")
    except sut.ThemeInvalidError:
        print("PASS: test_theme_invalid_route_fail_loud")


def test_theme_valid_route_pass():
    """合法主题直通"""
    tid, note = sut.resolve_theme({"style_theme": "kan-shi"}, {"themes": {"paper-ink": {}, "kan-shi": {}}})
    assert tid == "kan-shi" and note is None
    print("PASS: test_theme_valid_route_pass")


def test_theme_cli_override_semantics():
    """CLI --theme 合法值优先于 route 非法声明（compose 调用点语义）"""
    aesthetics = {"themes": {"paper-ink": {}, "kan-shi": {}}}
    theme = "kan-shi"  # CLI 显式
    valid_themes = set(aesthetics.get("themes", {}).keys())
    assert theme in valid_themes  # 直用，不进 resolve_theme
    try:
        sut.resolve_theme({"style_theme": "not-a-theme"}, aesthetics)
        raise SystemExit("FAIL: 无 CLI 覆盖时 route 非法仍应 raise")
    except sut.ThemeInvalidError:
        print("PASS: test_theme_cli_override_semantics")


if __name__ == "__main__":
    tests = [
        test_yaml_parser_basic,
        test_yaml_parser_brackets,
        test_validate_route_exit_not_ppt,
        test_validate_route_conversion_subset,
        test_validate_route_missing_source,
        test_ratio_auto_mode,
        test_ratio_explicit,
        test_ratio_ambiguous,
        test_parse_pol_content,
        test_deck_spec_basic,
        test_deck_spec_quote_and_ratio,
        test_deck_spec_rejects_bad_ratio,
        test_cjk_spacing,
        test_main_bullet_detection,
        test_build_pptx_smoke,
        test_verify_gate_catches_corrupt,
        test_theme_invalid_route_fail_loud,
        test_theme_valid_route_pass,
        test_theme_cli_override_semantics,
    ]
    passed = failed = 0
    for t in tests:
        try:
            t()
            passed += 1
        except Exception as e:
            print(f"FAIL: {t.__name__}: {e}")
            failed += 1
    print(f"\n{'='*40}")
    print(f"结果: {passed} PASS, {failed} FAIL ({len(tests)} total)")
    sys.exit(0 if failed == 0 else 1)
